"""
docx_service.py  Server-side Word (.docx) generation for 7G House.

Builds a branded A4 Word document using python-docx with the same content
as the PDF report produced by pdf_service.py:

  • Title block with project specs
  • Description blurb
  • Cost estimate section (breakdown table + totals + notes)
  • Room dimensions table
  • One project image per section, full-width, with label caption

Supports English and French (lang="en" | "fr").

Images are fetched directly by the backend to avoid browser CORS issues.
"""

import io
import os
import logging
import asyncio
from typing import Optional

import httpx
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.pdf_service import _t

logger = logging.getLogger(__name__)

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

# ── Colour palette (matches pdf_service) ──────────────────────────────────────
CLR_PRIMARY = RGBColor(0x1A, 0x52, 0x76)   # #1a5276
CLR_ACCENT  = RGBColor(0x2E, 0x86, 0xC1)   # #2e86c1
CLR_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
CLR_DARK    = RGBColor(0x1A, 0x1A, 0x2E)   # #1a1a2e
CLR_MUTED   = RGBColor(0x66, 0x66, 0x66)   # #666666

_HEX_PRIMARY = "1A5276"
_HEX_ACCENT  = "2E86C1"
_HEX_BG      = "F8F9FA"
_HEX_WHITE   = "FFFFFF"
_HEX_DARK    = "1A1A2E"
_HEX_MUTED   = "666666"
_HEX_BORDER  = "D0D0D0"

FONT_NAME = "Noto Sans"


def _set_run(run, size: int = 11, bold: bool = False, color: Optional[RGBColor] = None,
             italic: bool = False) -> None:
    """Style a run with brand font, size, weight and colour."""
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), FONT_NAME)
    rfonts.set(qn("w:hAnsi"), FONT_NAME)
    rfonts.set(qn("w:eastAsia"), FONT_NAME)


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _cell_text(cell, text: str, size: int = 11, bold: bool = False,
               color: Optional[RGBColor] = None) -> None:
    cell.paragraphs[0].text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    _set_run(p.add_run(text), size=size, bold=bold, color=color)


def _row_background(table, row_idx: int, fill: str) -> None:
    for cell in table.rows[row_idx].cells:
        _shade_cell(cell, fill)


def _add_page_number_footer(section, page_num_label: str, tagline: str) -> None:
    """Footer with tagline (left) and dynamic PAGE field (right)."""
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.5), WD_ALIGN_PARAGRAPH.RIGHT)

    _set_run(p.add_run(tagline), size=9, color=CLR_MUTED)
    p.add_run("\t")
    _set_run(p.add_run(page_num_label + " "), size=9, color=CLR_MUTED)

    run = p.add_run()
    run.font.name = FONT_NAME
    run.font.size = Pt(9)
    run.font.color.rgb = CLR_MUTED
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def _new_page(doc: Document) -> None:
    doc.add_page_break()


def _add_spec_section(doc: Document, generation_data: dict, lang: str) -> None:
    gen_type = generation_data.get("generation_type", "house_plan")
    spec = generation_data.get("floor_plan_spec") or {}

    labels = {
        "region": _t("region", lang),
        "division": _t("division", lang),
        "bedrooms": _t("bedrooms", lang),
        "bathrooms": _t("bathrooms", lang),
        "gross_area": _t("gross_area", lang),
        "kitchen": _t("kitchen", lang),
        "extras": _t("extras", lang),
        "style": _t("style", lang),
        "roof": _t("roof", lang),
        "foundation": _t("foundation", lang),
        "key_rooms": _t("key_rooms", lang),
        "outdoor": _t("outdoor", lang),
    }

    rows: list[tuple[str, str]] = []
    if gen_type == "floor_plan":
        rows = [
            (labels["region"], spec.get("region", "")),
            (labels["division"], spec.get("division", "")),
            (labels["bedrooms"], str(spec.get("num_bedrooms", ""))),
            (labels["bathrooms"], str(spec.get("num_bathrooms", ""))),
            (labels["gross_area"], f"{spec.get('gross_area', '')} m²"),
            (labels["kitchen"], str(spec.get("kitchen_type", "")).capitalize()),
        ]
        if extras := spec.get("extras"):
            rows.append((labels["extras"], ", ".join(extras)))
    else:
        rows = [
            (labels["style"], generation_data.get("house_style", "")),
            (labels["region"], generation_data.get("region", "")),
            (labels["division"], generation_data.get("division", "")),
            (labels["bedrooms"], str(generation_data.get("num_bedrooms", ""))),
            (labels["bathrooms"], str(generation_data.get("num_bathrooms", ""))),
            (labels["gross_area"], f"{generation_data.get('gross_area', '')} m²"),
            (labels["roof"], generation_data.get("roof_type", "")),
            (labels["foundation"], generation_data.get("foundation", "")),
        ]
        if ks := generation_data.get("key_rooms"):
            rows.append((labels["key_rooms"], ", ".join(ks)))
        if os_ := generation_data.get("outdoor_spaces"):
            rows.append((labels["outdoor"], ", ".join(os_)))

    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, (label, value) in enumerate(rows):
        _cell_text(table.rows[i].cells[0], label.upper(), size=10, bold=True, color=CLR_MUTED)
        _cell_text(table.rows[i].cells[1], value, size=11, bold=True, color=CLR_PRIMARY)
        if i % 2 == 0:
            _row_background(table, i, _HEX_BG)


def _add_cost_section(doc: Document, cost: dict, lang: str) -> None:
    hub_name = cost.get("hub_name", "")
    distance_km = cost.get("distance_km", 0)
    grand_total = cost.get("grand_total_fcfa", 0)
    grand_total_usd = cost.get("grand_total_usd", 0)
    base_per_m2 = cost.get("base_cost_per_m2", 0)
    breakdown = cost.get("breakdown", [])
    notes = cost.get("notes", [])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p.add_run(_t("cost_estimate_title", lang)), size=22, bold=True, color=CLR_PRIMARY)

    hub_line = f"{hub_name}"
    if distance_km > 0:
        hub_line += f"  {distance_km:.0f} km {_t('cost_hub_distance', lang)}"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p.add_run(hub_line), size=11, color=CLR_MUTED)

    labels_map = {
        "Materials": _t("cost_materials", lang),
        "Transport": _t("cost_transport", lang),
        "Labour": _t("cost_labour", lang),
        "Terrain Adjustment": _t("cost_terrain", lang),
        "Contingency": _t("cost_contingency", lang),
    }

    table = doc.add_table(rows=1 + len(breakdown) + 1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    header = table.rows[0].cells
    _cell_text(header[0], _t("cost_component", lang).upper(), size=11, bold=True, color=CLR_WHITE)
    _cell_text(header[1], _t("cost_ratio", lang).upper(), size=11, bold=True, color=CLR_WHITE)
    _cell_text(header[2], _t("cost_fcfa", lang).upper(), size=11, bold=True, color=CLR_WHITE)
    _row_background(table, 0, _HEX_PRIMARY)

    for i, item in enumerate(breakdown, start=1):
        label = labels_map.get(item["label"], item["label"])
        cells = table.rows[i].cells
        _cell_text(cells[0], label, size=11, bold=True, color=CLR_DARK)
        _cell_text(cells[1], f'{item["pct"]}%', size=10, color=CLR_MUTED)
        _cell_text(cells[2], f'{item["amount"]:,}', size=11, bold=True, color=CLR_PRIMARY)
        if i % 2 == 0:
            _row_background(table, i, _HEX_BG)

    total_cells = table.rows[len(breakdown) + 1].cells
    _cell_text(total_cells[0], _t("cost_grand_total", lang), size=12, bold=True, color=CLR_WHITE)
    _cell_text(total_cells[1], "", size=11, color=CLR_WHITE)
    _cell_text(total_cells[2], f"{grand_total:,} FCFA", size=12, bold=True, color=CLR_WHITE)
    _row_background(table, len(breakdown) + 1, _HEX_PRIMARY)

    info = (
        f'{_t("cost_usd_equiv", lang)}: ~{grand_total_usd:,} USD  |  '
        f'{_t("cost_per_m2", lang)}: {base_per_m2:,} FCFA/m²'
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p.add_run(info), size=10, color=CLR_MUTED)

    if notes:
        p = doc.add_paragraph()
        _set_run(p.add_run(_t("cost_notes", lang)), size=11, bold=True, color=CLR_DARK)
        for note in notes:
            p = doc.add_paragraph(style="List Bullet")
            _set_run(p.add_run(note), size=10, color=CLR_MUTED)


def _add_rooms_section(doc: Document, rooms: list[dict], lang: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p.add_run(_t("rooms_title", lang)), size=22, bold=True, color=CLR_PRIMARY)

    total_area = sum(r.get("area_m2", 0) for r in rooms)
    avg = total_area / len(rooms) if rooms else 0
    summary = (
        f'{len(rooms)} {_t("rooms_total", lang)}  |  '
        f'{_t("rooms_total_area", lang)}: {total_area:.1f} m²  |  '
        f'{_t("rooms_avg_size", lang)}: {avg:.1f} m²'
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p.add_run(summary), size=11, color=CLR_MUTED)

    table = doc.add_table(rows=1 + len(rooms) + 1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    header = table.rows[0].cells
    headers = [
        _t("rooms_col_room", lang),
        _t("rooms_col_code", lang),
        _t("rooms_col_width", lang),
        _t("rooms_col_height", lang),
        _t("rooms_col_area", lang),
    ]
    for j, htext in enumerate(headers):
        _cell_text(header[j], htext.upper(), size=10.5, bold=True, color=CLR_WHITE)
    _row_background(table, 0, _HEX_PRIMARY)

    for i, room in enumerate(rooms, start=1):
        cells = table.rows[i].cells
        _cell_text(cells[0], room.get("name", ""), size=10.5, bold=True, color=CLR_DARK)
        _cell_text(cells[1], str(room.get("code", "")), size=10, color=CLR_MUTED)
        _cell_text(cells[2], f'{room.get("width_m", 0):.1f}', size=10.5, bold=True, color=CLR_PRIMARY)
        _cell_text(cells[3], f'{room.get("height_m", 0):.1f}', size=10.5, bold=True, color=CLR_PRIMARY)
        _cell_text(cells[4], f'{room.get("area_m2", 0):.1f}', size=10.5, bold=True, color=CLR_PRIMARY)
        if i % 2 == 0:
            _row_background(table, i, _HEX_BG)

    total_cells = table.rows[len(rooms) + 1].cells
    _cell_text(total_cells[0], _t("rooms_total", lang), size=11, bold=True, color=CLR_WHITE)
    _cell_text(total_cells[1], "", size=10, color=CLR_WHITE)
    _cell_text(total_cells[2], "", size=10, color=CLR_WHITE)
    _cell_text(total_cells[3], "", size=10, color=CLR_WHITE)
    _cell_text(total_cells[4], f"{total_area:.1f}", size=11, bold=True, color=CLR_WHITE)
    _row_background(table, len(rooms) + 1, _HEX_PRIMARY)


def _assemble_docx(generation_data: dict, downloaded: list[tuple[str, bytes]],
                   lang: str = "en") -> Optional[io.BytesIO]:
    """Build the Word document (runs in a worker thread)."""
    try:
        doc = Document()
        section = doc.sections[0]
        section.page_width = Cm(21.0)      # A4
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.6)
        section.right_margin = Cm(1.6)

        normal = doc.styles["Normal"]
        normal.font.name = FONT_NAME
        normal.font.size = Pt(11)

        # ── Title block ──────────────────────────────────────────────────────
        gen_type = generation_data.get("generation_type", "house_plan")
        spec = generation_data.get("floor_plan_spec") or {}

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(60)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run(p.add_run("7G HOUSE"), size=30, bold=True, color=CLR_PRIMARY)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run(p.add_run(_t("header_subtitle", lang).upper()), size=12, bold=True,
                 color=CLR_ACCENT)

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(40)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run(p.add_run(_t("project_report", lang)), size=26, bold=True, color=CLR_PRIMARY)

        if gen_type == "floor_plan":
            bedrooms = spec.get("num_bedrooms", 0)
            title = _t("floor_plan_title", lang).format(bedrooms=bedrooms)
        else:
            style = generation_data.get("house_style") or "Residential"
            title = _t("house_design_title", lang).format(style=style)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run(p.add_run(title), size=15, color=CLR_DARK)

        client = generation_data.get("client_name")
        if client:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_run(p.add_run(_t("prepared_for", lang).format(client=client)),
                     size=11, color=CLR_MUTED)

        # ── Spec table ───────────────────────────────────────────────────────
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(24)
        _add_spec_section(doc, generation_data, lang)

        # ── Description blurb ────────────────────────────────────────────────
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        _set_run(p.add_run(_t("desc_blurb", lang)), size=10, italic=True, color=CLR_MUTED)

        # ── Cost estimate (if available) ─────────────────────────────────────
        cost_estimate = generation_data.get("cost_estimate")
        if cost_estimate:
            _new_page(doc)
            _add_cost_section(doc, cost_estimate, lang)

        # ── Room dimensions (if available) ───────────────────────────────────
        room_measurements = generation_data.get("room_measurements", [])
        if room_measurements:
            _new_page(doc)
            _add_rooms_section(doc, room_measurements, lang)

        # ── Project images ───────────────────────────────────────────────────
        for img_index, (label, img_bytes) in enumerate(downloaded, start=1):
            _new_page(doc)
            try:
                stream = io.BytesIO(img_bytes)
                doc.add_picture(stream, width=Inches(6.9))
                last_paragraph = doc.paragraphs[-1]
                last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception as e:
                logger.warning(f"docx_service: cannot embed image {img_index}: {e}")
                continue

            cap = label or f"{_t('view_label', lang)} {img_index}"
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            _set_run(p.add_run(cap.upper()), size=11, bold=True, color=CLR_MUTED)

        # ── Footer with page number ─────────────────────────────────────────
        _add_page_number_footer(section, _t("page", lang), _t("footer_tagline", lang))

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        buf.name = "Project_Report.docx"
        return buf
    except Exception as e:
        logger.error(f"docx_service: failed to build document: {e}")
        return None


async def build_project_docx(generation_data: dict, lang: str = "en") -> Optional[io.BytesIO]:
    """
    Build a branded Word document from a generation document.
    Returns a BytesIO buffer with name set to 'Project_Report.docx',
    or None on failure.
    """
    images: list[dict] = generation_data.get("images", [])
    headers = {"User-Agent": "7G-House-DOCX/1.0"}

    downloaded: list[tuple[str, bytes]] = []
    async with httpx.AsyncClient(headers=headers, timeout=30.0) as client:
        for img_meta in images:
            url = img_meta.get("url", "")
            label = img_meta.get("label", "")
            if not url:
                continue
            try:
                resp = await client.get(url, follow_redirects=True)
                resp.raise_for_status()
                downloaded.append((label, resp.content))
            except Exception as e:
                logger.warning(f"docx_service: failed to download image '{url[:60]}': {e}")

    try:
        buf = await asyncio.to_thread(_assemble_docx, generation_data, downloaded, lang)
    except Exception as e:
        logger.error(f"docx_service: failed to assemble document: {e}")
        return None

    if buf is None:
        logger.error("docx_service: no document generated")
        return None

    logger.info(f"docx_service: built document with {len(downloaded)} images ({len(buf.getvalue())} bytes)")
    return buf


async def upload_docx_to_storage(docx_buf: io.BytesIO, generation_id: str) -> Optional[str]:
    """Upload the DOCX buffer to Firebase Storage and return the public URL."""
    from app.db.firebase import bucket
    import uuid

    if bucket is None:
        logger.error("docx_service: Firebase Storage bucket not initialised")
        return None

    try:
        blob_path = f"project_docx/{generation_id}/{uuid.uuid4().hex}.docx"
        blob = bucket.blob(blob_path)
        docx_buf.seek(0)
        blob.upload_from_file(docx_buf, content_type=DOCX_MEDIA_TYPE)
        blob.make_public()
        url = blob.public_url
        logger.info(f"docx_service: uploaded document → {url}")
        return url
    except Exception as e:
        logger.error(f"docx_service: upload failed: {e}")
        return None


async def delete_docx_from_storage(docx_url: str) -> bool:
    """Delete a DOCX file from Firebase Storage given its public URL."""
    from app.db.firebase import bucket

    if bucket is None:
        logger.error("docx_service: Firebase Storage bucket not initialised")
        return False

    if not docx_url:
        return False

    try:
        from urllib.parse import urlparse
        parsed = urlparse(docx_url)
        blob_path = parsed.path.lstrip("/")
        if blob_path.startswith(f"{bucket.name}/"):
            blob_path = blob_path[len(bucket.name) + 1:]
        blob = bucket.blob(blob_path)
        if blob.exists():
            blob.delete()
            logger.info(f"docx_service: deleted document from storage → {blob_path}")
            return True
        else:
            logger.warning(f"docx_service: document blob not found → {blob_path}")
            return False
    except Exception as e:
        logger.error(f"docx_service: delete failed: {e}")
        return False